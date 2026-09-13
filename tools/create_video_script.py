from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = r"E:\movie-discovery-app\ReelScout Project Walkthrough Script.docx"


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    paragraph.paragraph_format.space_after = Pt(3)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("ReelScout Movie Discovery App")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Long Form Project Walkthrough Video Script")
run.bold = True
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(48, 76, 108)

intro = doc.add_paragraph()
intro.add_run("Purpose. ").bold = True
intro.add_run(
    "This script is designed for a ten to twelve minute project demonstration. "
    "Read it naturally rather than trying to recite every word exactly. The bracketed "
    "directions describe what should be visible on screen while the narration is spoken."
)

table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
headers = ["Video length", "Recommended delivery"]
for cell, text in zip(table.rows[0].cells, headers):
    shade_cell(cell, "244F70")
    set_cell_text(cell, text, bold=True, color="FFFFFF")
for left, right in [
    ("10 to 12 minutes", "Calm, clear, and confident"),
    ("Screen recording", "Show the running app and a few code files"),
    ("Main goal", "Demonstrate product thinking and technical ownership"),
]:
    cells = table.add_row().cells
    set_cell_text(cells[0], left)
    set_cell_text(cells[1], right)

segments = [
    (
        "Opening and Project Context",
        "0:00 to 0:50",
        "[Show the ReelScout home page. Start on the desktop layout, then briefly resize the browser to a phone width.]\n\n"
        "Hello, and welcome to my walkthrough of ReelScout, a full-stack movie discovery application. "
        "I built this project as a product rather than as a simple API demonstration. The objective is to "
        "help a user discover something interesting immediately, search for a specific title when they have "
        "one in mind, inspect useful details, and save movies for later.\n\n"
        "The application has a React and TypeScript frontend, a Node.js and Express backend, and a movie data "
        "provider integration. The browser communicates with my backend instead of calling the external movie "
        "service directly. That gives me a clean place to protect credentials, normalize provider data, cache "
        "repeated requests, and handle failures consistently. In this video I will demonstrate the main user "
        "flows, then explain the important technical decisions behind them."
    ),
    (
        "Product Experience and Discovery",
        "0:50 to 2:05",
        "[Show the home page from the top. Scroll slowly through the movie grid.]\n\n"
        "The first screen is intentionally focused on discovery. A user does not need to know a movie title "
        "before using the app. The page introduces the product with a simple search field and then presents a "
        "movie grid. Each movie card gives the user the title, year, rating, poster, and a save action.\n\n"
        "The layout is built to remain useful as the number of results grows. Cards use a responsive grid, "
        "posters preserve a consistent aspect ratio, and long titles are truncated so one unusually long title "
        "does not stretch the whole page. If a provider record does not include a poster, the interface uses a "
        "fallback rather than leaving a broken image.\n\n"
        "The two controls on the discovery page allow the user to filter by genre and change the ordering. For "
        "example, the user can view all genres, select action or drama, and sort by popularity, rating, or release "
        "date. When a control changes, the page requests the appropriate result set and resets pagination so the "
        "user does not accidentally remain on an invalid later page."
    ),
    (
        "Search and Result States",
        "2:05 to 3:10",
        "[Search for a recognizable title, then search for a phrase that produces no results.]\n\n"
        "The search flow is kept separate from discovery because the user intent is different. Discovery is broad "
        "and exploratory, while search is direct. From the search box I can enter a movie title and navigate to a "
        "search URL that contains the query. This makes the result page refreshable and shareable.\n\n"
        "The interface has an explicit loading state while the request is in progress. Instead of leaving an empty "
        "area on the page, it displays skeleton cards that communicate the shape of the result. If the query returns "
        "nothing, the user sees a useful empty state with a suggestion to try another phrase. If the request fails, "
        "the error state includes a retry action. These states are important because a real external API can be slow, "
        "temporarily unavailable, or return no matching records.\n\n"
        "On the backend, the search query is validated before it is sent to the provider. Queries shorter than two "
        "characters are rejected, and the page parameter is constrained to a safe range. This prevents malformed "
        "requests from reaching the provider and gives the frontend a predictable response shape."
    ),
    (
        "Movie Details and Navigation",
        "3:10 to 4:05",
        "[Open a movie card and show the details page. Point to the poster, rating, genres, and save button.]\n\n"
        "Selecting a movie opens a dedicated details route. The details page shows the larger poster, title, overview, "
        "rating, release year, runtime when available, and genre tags. This screen is intentionally more spacious than "
        "the grid because it gives the user room to decide whether the movie belongs on their list.\n\n"
        "The details route uses the movie ID, and the backend fetches the complete movie record. The frontend does not "
        "need to understand the provider's raw field names. It consumes the application's own Movie type, which includes "
        "safe values for missing descriptions, missing dates, missing ratings, and missing artwork.\n\n"
        "The back link takes the user back to discovery, while the main navigation always provides access to Discover and "
        "Wishlist. This keeps the three core contexts visible and makes the application easy to understand even for a "
        "first-time user."
    ),
    (
        "Wishlist Persistence",
        "4:05 to 5:05",
        "[Save a movie, navigate to Wishlist, refresh the page, and show that the item remains.]\n\n"
        "The wishlist is the main persistent feature in the assignment. From a movie card or the details page, I can "
        "save a movie. The button immediately communicates whether the title is saved, and the Wishlist page displays "
        "the saved collection using the same reusable movie card component.\n\n"
        "This version does not require account registration, so it uses a device identity. On the first visit, the browser "
        "generates a UUID and stores it in localStorage. Each wishlist request sends that value in the X-Device-Id header. "
        "The backend uses the device ID as the owner of the saved movies and persists the movie snapshot in a local JSON "
        "file. That means the wishlist survives page refreshes, browser restarts, and backend process restarts for the same "
        "device.\n\n"
        "The store prevents duplicates by indexing movies by ID within each device collection. Removing a movie updates the "
        "same persistent store. For a production product I would replace this anonymous device identity with authenticated "
        "accounts and a database such as PostgreSQL, but the current choice keeps the assignment simple while still showing "
        "real persistence and a clear upgrade path."
    ),
    (
        "Backend Architecture",
        "5:05 to 6:35",
        "[Switch to the backend folder in the editor. Show app.ts, movieProvider.ts, and wishlist.ts.]\n\n"
        "The backend is organized around a small set of responsibilities. The Express app owns routing, middleware, "
        "validation, and error conversion. The movie provider service owns communication with TMDB and maps external "
        "records into the internal Movie model. The cache service stores short-lived responses so repeated requests do "
        "not unnecessarily consume provider calls. The wishlist service owns device-based persistence.\n\n"
        "The provider boundary is important. External movie services can change field names, omit values, enforce rate "
        "limits, or return an error response. By handling those concerns inside the provider service, the React application "
        "remains independent of the provider. The backend exposes endpoints such as GET /api/movies/discover, GET "
        "/api/movies/search, GET /api/movies/:id, and the wishlist endpoints.\n\n"
        "The backend also uses Helmet for security headers, CORS for the frontend origin, JSON request limits, rate limiting, "
        "and Zod validation. Errors are returned using a consistent object with an error code and message. Provider failures "
        "become a controlled service-unavailable response instead of an unhandled exception. During local development the "
        "application can use a demo catalog fallback, which keeps the interface testable if the external provider or network "
        "is temporarily unavailable."
    ),
    (
        "External API Data Flow",
        "6:35 to 7:35",
        "[Show the browser Network tab, then show the backend terminal and the .env.example file without showing secrets.]\n\n"
        "The data flow for a discovery request has several deliberate steps. First, React requests my own endpoint, not TMDB. "
        "The backend validates the page, genre, and sort parameters. It then checks the in-memory cache using a key built from "
        "those inputs. If the same request was made recently, the cached normalized result can be returned immediately.\n\n"
        "If the cache does not contain a result, the provider service creates a TMDB request with the server-side API key, "
        "language, page, and filter parameters. The request has an AbortController timeout. A slow or unavailable provider "
        "does not leave the server hanging indefinitely.\n\n"
        "When the response arrives, the mapper converts fields such as poster_path and vote_average into posterUrl and rating. "
        "It also provides fallback text for incomplete data and converts provider genre IDs into readable names where possible. "
        "The normalized result is placed into the cache and returned to React with pagination metadata. The frontend can then "
        "render without knowing whether a result came from the cache, the provider, or the local demo catalog."
    ),
    (
        "Responsive Mobile Design",
        "7:35 to 8:35",
        "[Resize the browser through desktop, tablet, and narrow phone widths. Show the navigation, grid, filters, and detail page.]\n\n"
        "Responsive behavior was treated as part of the product rather than as a final visual adjustment. At desktop widths, "
        "the application uses the available space for a comfortable multi-column movie grid. At tablet and phone widths, the "
        "grid reduces its columns, the navigation becomes vertical, and the filter controls expand to use the available width.\n\n"
        "The search form is especially important on small screens. On a narrow phone it becomes a vertical control with a full-width "
        "button, which avoids a cramped input and gives the button a comfortable touch target. Buttons use a minimum height that "
        "is easier to tap, and the poster cards keep their aspect ratio so the layout does not jump as images load.\n\n"
        "The details page also changes from a two-column layout to a single column. The poster is centered, the metadata is readable, "
        "and the save button expands to the full content width. I also added safer heading sizes and spacing for very narrow screens. "
        "The result is a product that remains usable rather than simply shrinking the desktop layout."
    ),
    (
        "Testing and Verification",
        "8:35 to 9:30",
        "[Show the terminal running the test and build commands. Then show one backend endpoint response.]\n\n"
        "I verified the backend with automated tests for two important boundaries. The movie mapper test checks that incomplete provider "
        "data receives safe defaults and that genre IDs are normalized. The wishlist test checks that the same movie is not duplicated "
        "and that one device cannot see another device's items.\n\n"
        "The project also has separate build commands for the backend and frontend. The backend TypeScript compiler verifies the server "
        "types, while Vite creates the production frontend bundle. I can manually test the running API with the health endpoint, the "
        "discovery endpoint, the search endpoint, a movie details endpoint, and the wishlist endpoints.\n\n"
        "The manual test checklist includes refreshing pages, closing and reopening the browser, changing filters quickly, trying a query "
        "with no results, disabling the network, checking missing poster fallbacks, and viewing the app on a mobile viewport. These tests "
        "exercise the real-world scenarios described in the assignment rather than only testing the happy path."
    ),
    (
        "Technical Tradeoffs",
        "9:30 to 10:35",
        "[Show the README sections for architecture, decisions, limitations, and AI transparency.]\n\n"
        "There are several tradeoffs I made consciously. I selected React for the web instead of React Native because the assignment "
        "emphasizes responsive behavior across screen sizes and because a browser-based review is easy to run. I selected Express because "
        "it is small, familiar, and sufficient for a clear service boundary.\n\n"
        "I used a file-backed wishlist rather than adding authentication and a full relational database. That keeps setup simple for a hiring "
        "assignment and still demonstrates persistence. The code is organized so the storage implementation can be replaced without changing "
        "the frontend contract. In production, I would use authenticated users, PostgreSQL, and a shared cache such as Redis.\n\n"
        "I also chose to normalize provider data at the backend boundary rather than pass raw provider responses to the browser. This adds a small "
        "amount of code now, but it protects the product from provider-specific changes and keeps the UI components easier to maintain. Finally, "
        "I kept the demo fallback enabled for local development so a reviewer can run the app even if a provider key or external network is not "
        "available."
    ),
    (
        "What I Would Improve Next",
        "10:35 to 11:25",
        "[Show the wishlist page and then a short list of future improvements on screen.]\n\n"
        "With more time, I would improve the application in four areas. First, I would add real user accounts, secure session handling, and a "
        "PostgreSQL database so a wishlist could follow a user across devices. Second, I would add more discovery options, including release "
        "year ranges, language, watch providers, cast, and runtime.\n\n"
        "Third, I would add a richer testing layer. That would include frontend component tests, API integration tests using a mocked provider, and "
        "end-to-end browser tests for search, details, and wishlist flows. I would also add continuous integration so every pull request runs the "
        "same build and test checks.\n\n"
        "Fourth, I would prepare the application for production deployment with structured logging, a shared cache, monitoring, image optimization, "
        "and a managed database. I would also add a privacy policy and make sure the TMDB attribution requirements are visible in the deployed UI."
    ),
    (
        "Closing Summary",
        "11:25 to 12:00",
        "[Return to the home page and show a short sequence: discover, search, open details, save to wishlist.]\n\n"
        "To summarize, ReelScout is a full-stack movie discovery product with a responsive React interface, a Node.js abstraction layer, external "
        "movie data integration, caching, validation, error handling, and persistent wishlist behavior. The main design goal was to make the app "
        "feel useful even when the user has no specific movie in mind, while still supporting fast search and detailed inspection.\n\n"
        "The implementation also demonstrates how I think about real application conditions: incomplete third-party data, provider failures, repeated "
        "requests, large result sets, responsive layouts, and persistence. The code and README document the tradeoffs, and the project is structured "
        "so the local storage and provider implementations can grow into production versions. Thank you for watching."
    ),
]

for heading, timing, body in segments:
    doc.add_heading(heading, level=1)
    timing_run = doc.add_paragraph()
    timing_run.add_run(f"Estimated time: {timing}").bold = True
    for block in body.split("\n\n"):
        if block.startswith("[") and "]" in block:
            p = doc.add_paragraph()
            r = p.add_run(block)
            r.italic = True
            r.font.color.rgb = RGBColor(80, 80, 80)
        else:
            doc.add_paragraph(block)

doc.add_heading("Recording Checklist", level=1)
for item in [
    "Start both backend and frontend servers before recording.",
    "Do not show the TMDB API key or any .env file containing a secret.",
    "Use a browser width that makes the desktop layout clear, then demonstrate a phone width.",
    "Keep the browser Network tab closed unless you are intentionally demonstrating the backend boundary.",
    "Pause briefly after each major user action so the reviewer can follow the result.",
    "Run npm test and npm run build before recording the final take.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.core_properties.title = "ReelScout Movie Discovery App Project Walkthrough Video Script"
doc.core_properties.subject = "Long form video narration script for the movie discovery project"
doc.core_properties.author = ""
doc.save(OUTPUT)
print(OUTPUT)
